"""离线验证：使用临时目录，不连接邮箱，不修改实际收取记录。"""
import tempfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from 收作业 import export_excel, read_roster


def main():
    sid, name = '0123456789', '示例甲'
    def record(week):
        return dict(number=week, sid=sid, name=name, status='已交', files=['附件.pdf'],
                    subject=f'第{week}周作业', sender='', date='')
    state = {'third': record(3), 'fourth': record(4), 'duplicate': record(3)}
    with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as folder:
        roster = Path(folder) / '名单.xlsx'
        source = Workbook()
        source.active.append(['学号', '姓名'])
        source.active.append([sid, name])
        source.active.append(['1123456789', '示例乙'])
        source.save(roster)
        source.close()
        assert read_roster(roster) == {sid: name, '1123456789': '示例乙'}
        export_excel(Path(folder), state, roster, 4)
        wb = load_workbook(Path(folder) / '作业统计.xlsx')
        try:
            rows = list(wb['提交统计'].values)
            assert rows[0] == ('姓名', '学号', '第3周', '第4周', '已交周数', '备注')
            assert len(rows) == 3
            assert rows[1][0:5] == (name, sid, '已交', '已交', 2)
            assert rows[2][2:4] == (None, None)
            assert wb['邮件明细'].max_row == 4
        finally:
            wb.close()
        export_excel(Path(folder), state, roster, 6)
        export_excel(Path(folder), state, roster, 7)
        wb = load_workbook(Path(folder) / '作业统计.xlsx')
        try:
            assert list(wb['提交统计'].values)[0] == ('姓名', '学号', '第3周', '第4周', '第6周', '第7周', '已交周数', '备注')
        finally:
            wb.close()
    print('验证通过：名单完整、跨周保留、重复提交去重、未交单元格留空。')


if __name__ == '__main__':
    main()
