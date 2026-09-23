"""从 IMAP 邮箱下载作业附件并生成 Excel；运行 python 收作业.py --help。"""
import argparse
import getpass
import hashlib
import imaplib
import json
import os
import re
import ssl
import sys
import unicodedata
from datetime import date
from email import policy
from email.parser import BytesParser
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent / '本地依赖'))

from 学号范围 import in_scope


def parse_subject(subject):
    subject = unicodedata.normalize("NFKC", subject).strip()
    match = re.fullmatch(
        r"第\s*([0-9一二三四五六七八九十百零〇两]+)\s*周作业\s*[+_\-\s]*([0-9]{10})[+_\-\s]*([^+\r\n]+)", subject
    )
    if not match:
        return None
    number, sid, name = match.groups()
    if number.isascii() and number.isdigit():
        number = int(number)
    else:
        digits = dict(zip("零〇一二两三四五六七八九", [0, 0, 1, 2, 2, 3, 4, 5, 6, 7, 8, 9]))
        total = pending = 0
        for char in number:
            if char in digits:
                pending = digits[char]
            else:
                total += (pending or 1) * {"十": 10, "百": 100}[char]
                pending = 0
        number = total + pending
    name = name.strip()
    return (number, sid, name) if number > 0 and name else None


def safe_name(value):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', value).strip(' .')[:100]
    return value or "未命名"


def save_json(path, value):
    temporary = path.with_suffix('.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    temporary.replace(path)


def parse_since(value):
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        raise argparse.ArgumentTypeError('日期格式须为 YYYY-MM-DD，例如 2026-09-01')
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError('日期无效，请使用 YYYY-MM-DD，例如 2026-09-01') from exc


def collect(client, args, root, state):
    status, _ = client.select(args.mailbox, readonly=True)
    if status != 'OK':
        raise RuntimeError('无法打开邮箱文件夹')
    _, validity = client.response('UIDVALIDITY')
    if not validity or not validity[0]:
        raise RuntimeError('邮箱未返回 UIDVALIDITY，无法安全识别邮件')
    namespace = '|'.join([args.server, args.email, args.mailbox, validity[0].decode()])
    if args.since:
        months = ('Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec')
        since = f'{args.since.day:02d}-{months[args.since.month - 1]}-{args.since.year:04d}'
        status, data = client.uid('search', None, 'SINCE', since)
    else:
        status, data = client.uid('search', None, 'ALL')
    if status != 'OK':
        raise RuntimeError('搜索邮件失败')
    for uid in data[0].split():
        key = namespace + '|' + uid.decode()
        try:
            status, response = client.uid('fetch', uid, '(BODY.PEEK[])')
            if status != 'OK':
                raise RuntimeError('读取邮件失败')
            raw = b''.join(item[1] for item in response if isinstance(item, tuple))
            if not raw:
                raise RuntimeError('邮件内容为空')
            msg = BytesParser(policy=policy.default).parsebytes(raw)
            subject = str(msg.get('Subject', ''))
            parsed = parse_subject(subject)
            if parsed and not in_scope(parsed[1]):
                continue
            if not parsed and '作业' not in subject:
                continue
            if parsed and args.assignment and parsed[0] != args.assignment:
                continue
            record = dict(subject=subject, sender=str(msg.get('From', '')),
                          date=str(msg.get('Date', '')), number=None, sid='', name='', files=[], status='主题格式错误')
            if parsed:
                number, sid, name = parsed
                record.update(number=number, sid=sid, name=name, status='无附件')
                folder = root / f'第{number}周作业' / safe_name(f'{sid}_{name}')
                # 使用邮件唯一标识和附件序号区分重交及同名附件。
                token = hashlib.sha256(key.encode()).hexdigest()[:20]
                for index, part in enumerate(msg.walk(), 1):
                    if part.is_multipart():
                        continue
                    filename = part.get_filename()
                    if not filename and part.get_content_disposition() != 'attachment':
                        continue
                    payload = part.get_payload(decode=True)
                    if payload is None:
                        continue
                    folder.mkdir(parents=True, exist_ok=True)
                    target = folder / f'{token}_{index}_{safe_name(filename or "附件.bin")}'
                    if not target.exists() or target.read_bytes() != payload:
                        temporary = target.with_suffix(target.suffix + '.tmp')
                        temporary.write_bytes(payload)
                        temporary.replace(target)
                    record['files'].append(str(target.relative_to(root)))
                if record['files']:
                    record['status'] = '已交'
            state[key] = record
            save_json(root / '收取记录.json', state)
            print(f"[{record['status']}] {subject}")
        except (OSError, ValueError, RuntimeError, imaplib.IMAP4.error) as exc:
            print(f'[失败] 邮件 UID {uid.decode()}：{exc}；下次运行会重试')


def read_roster(roster):
    if Path(roster).suffix.lower() == '.xls':
        import xlrd
        source = xlrd.open_workbook(str(roster))
        try:
            sheet = source.sheet_by_index(0)
            rows = [sheet.row_values(i) for i in range(sheet.nrows)]
        finally:
            source.release_resources()
    else:
        from openpyxl import load_workbook
        source = load_workbook(roster, read_only=True, data_only=True)
        try:
            rows = list(source.active.values)
        finally:
            source.close()
    header_index = next((i for i, row in enumerate(rows)
                         if {'学号', '姓名'} <= {str(v).strip() for v in row}), None)
    if header_index is None:
        raise ValueError('名单中未找到“学号”和“姓名”表头')
    headers = [str(v).strip() for v in rows[header_index]]
    id_col, name_col = headers.index('学号'), headers.index('姓名')
    students = {}
    for row in rows[header_index + 1:]:
        value = row[id_col]
        if value is None or str(value).strip() == '':
            continue
        sid = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()
        if not re.fullmatch(r'\d{10}', sid):
            raise ValueError(f'名单中的学号必须是10位：{sid}')
        if not in_scope(sid):
            continue
        name = str(row[name_col] or '').strip()
        if sid in students and students[sid] != name:
            raise ValueError(f'名单中同一学号对应多个姓名：{sid}')
        students[sid] = name
    return students


def export_excel(root, state, roster=None, assignment=None):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    records = [r for r in state.values() if not r['sid'] or in_scope(r['sid'])]
    students = read_roster(roster) if roster else {}
    roster_ids = set(students)
    known_weeks = {r['number'] for r in records if r['number']} | ({assignment} if assignment else set())
    # 保留之前已建立的周列，包括已布置但尚无人提交的周次。
    existing_path = root / '作业统计.xlsx'
    if existing_path.exists():
        previous = load_workbook(existing_path, read_only=True, data_only=True)
        try:
            for cell in next(previous['提交统计'].iter_rows(max_row=1)):
                match = re.fullmatch(r'第(\d+)周(?:作业)?', str(cell.value))
                if match:
                    known_weeks.add(int(match.group(1)))
        finally:
            previous.close()
    numbers = sorted(known_weeks)
    for r in records:
        if r['sid']:
            students.setdefault(r['sid'], r['name'])
    wb = Workbook()
    summary = wb.active
    summary.title = '提交统计'
    summary.append(['姓名', '学号'] + [f'第{n}周' for n in numbers] + ['已交周数', '备注'])
    submitted = {(r['sid'], r['number']) for r in records if r['status'] == '已交'}
    for sid, name in students.items():
        marks = ['已交' if (sid, n) in submitted else None for n in numbers]
        notes = []
        if roster and sid not in roster_ids:
            notes.append('不在名单中')
        if any(r['sid'] == sid and r['name'] != name for r in records):
            notes.append('邮件姓名与名单或其他提交不一致')
        summary.append([name, sid] + marks + [marks.count('已交'), '；'.join(notes)])
    detail = wb.create_sheet('邮件明细')
    detail.append(['作业周次', '学号', '姓名', '状态', '邮件时间', '发件人', '主题', '附件数量', '附件相对路径'])
    for r in records:
        detail.append([r['number'], r['sid'], r['name'], r['status'], r['date'], r['sender'], r['subject'], len(r['files']), '\n'.join(r['files'])])
    for ws in wb:
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='4472C4')
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(60, max(16, max(len(str(c.value or '')) for c in column) + 2))
    destination = root / '作业统计.xlsx'
    temporary = root / '作业统计.tmp.xlsx'
    wb.save(temporary)
    temporary.replace(destination)
    print(f'统计已保存：{destination}')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--email', default=os.getenv('HOMEWORK_EMAIL'), help='收作业邮箱，也可设置 HOMEWORK_EMAIL')
    parser.add_argument('--server', default='imap.qq.com')
    parser.add_argument('--port', type=int, default=993)
    parser.add_argument('--mailbox', default='INBOX', help='默认收件箱')
    parser.add_argument('--output', default='作业收集')
    parser.add_argument('--roster', help='含学号、姓名的 .xls/.xlsx 名单；默认查找脚本旁唯一的“选课名单”文件')
    parser.add_argument('--assignment', type=int, help='只下载指定周次；Excel 始终汇总本目录所有周次')
    parser.add_argument('--since', type=parse_since, help='仅下载此收件日期及之后的邮件附件，格式 YYYY-MM-DD（含当天）；不清除本地历史记录')
    parser.add_argument('--export-only', action='store_true', help='不连接邮箱，只根据本地收取记录重新生成 Excel')
    parser.add_argument('--web-output', default='site', help='查询网页输出目录，默认 site，供 push 后发布')
    parser.add_argument('--title', default='高等数学 B', help='网页课程名称')
    args = parser.parse_args()
    if args.assignment is not None and args.assignment <= 0:
        parser.error('--assignment 必须大于0')
    try:
        import openpyxl  # 下载之前检查依赖
        if not args.roster:
            candidates = [p for p in Path(__file__).resolve().parent.glob('*选课名单*')
                          if p.suffix.lower() in ('.xls', '.xlsx') and not p.name.startswith('~$')]
            if len(candidates) > 1:
                raise ValueError('找到多份选课名单，请用 --roster 指定')
            if candidates:
                args.roster = candidates[0]
        if args.roster:
            read_roster(args.roster)  # 登录和下载前验证名单
        root = Path(args.output).resolve()
        root.mkdir(parents=True, exist_ok=True)
        state_path = root / '收取记录.json'
        state = json.loads(state_path.read_text(encoding='utf-8')) if state_path.exists() else {}
        if not args.export_only:
            args.email = args.email or input('收作业邮箱：').strip()
            code = os.getenv('HOMEWORK_AUTH_CODE') or getpass.getpass('邮箱 IMAP 授权码（输入不显示）：')
            with imaplib.IMAP4_SSL(args.server, args.port, ssl_context=ssl.create_default_context(), timeout=60) as client:
                client.login(args.email, code)
                collect(client, args, root, state)
        export_excel(root, state, args.roster, args.assignment)
        from 生成网页 import build_site
        build_site(root / '作业统计.xlsx', args.web_output, args.title)
    except ImportError:
        parser.exit(1, '请先安装依赖：python -m pip install -r requirements.txt\n')
    except (OSError, ValueError, KeyError, StopIteration, imaplib.IMAP4.error) as exc:
        parser.exit(1, f'运行失败：{exc}\n若 Excel 正在打开，请关闭后重试；已保存的附件和记录会保留。\n')


if __name__ == '__main__':
    main()
