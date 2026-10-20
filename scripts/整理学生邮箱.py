"""按规范主题更新学生邮箱；每个学生仅保留最新邮件的发件地址。"""
import argparse
from email.utils import getaddresses
import json
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook


def organize(root, records=None):
    from 收作业 import parse_subject, submission_order, in_scope

    root = Path(root)
    if records is None:
        records = json.loads((root / '收取记录.json').read_text(encoding='utf-8'))
    target = root / '学生邮箱.xlsx'
    wb = load_workbook(target) if target.exists() else Workbook()
    try:
        if not target.exists():
            wb.active.title = '学生邮箱'
            wb.active.append(['学号', '姓名', '邮箱（多个用分号分隔）', '备注'])
        ws = wb['学生邮箱']
        columns = {cell.value: cell.column for cell in ws[1]}
        rows = {str(ws.cell(row, columns['学号']).value): row for row in range(2, ws.max_row + 1)}
        chosen = {}
        for key in sorted(records, key=lambda key: submission_order(key, records[key])):
            record = records[key]
            parsed = parse_subject(record.get('subject', ''))
            if not parsed or not in_scope(parsed[1]):
                continue
            addresses = {addr.strip().lower() for _, addr in getaddresses([record.get('sender', '')])
                         if re.fullmatch(r'[^\s@<>]+@[^\s@<>]+\.[^\s@<>]+', addr.strip())}
            if len(addresses) == 1:
                chosen[parsed[1]] = (parsed[2], next(iter(addresses)))
        for sid, (name, address) in chosen.items():
            row = rows.get(sid, ws.max_row + 1)
            for header, value in [('学号', sid), ('姓名', name), ('邮箱（多个用分号分隔）', address),
                                  ('备注', ''), ('邮箱数量', 1), ('匹配邮件数', None)]:
                if header in columns:
                    cell = ws.cell(row, columns[header], value)
                    cell.value = value
                    if isinstance(value, str):
                        cell.data_type = 's'
        # 旧版的汇总明细不再维护，避免与当前单一邮箱冲突。
        for title in ('来源明细', '待人工核对'):
            if title in wb.sheetnames:
                del wb[title]
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        for header, width in [('学号', 18), ('姓名', 16), ('邮箱（多个用分号分隔）', 40)]:
            ws.column_dimensions[ws.cell(1, columns[header]).column_letter].width = width
        temporary = target.with_name('学生邮箱.tmp.xlsx')
        wb.save(temporary)
    finally:
        wb.close()
    temporary.replace(target)
    print(f'已更新 {len(chosen)} 位学生的邮箱：{target}')
    return target


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', default='作业收集_2026秋', help='本地作业数据目录')
    args = parser.parse_args()
    organize(args.output)
