"""从 IMAP 邮箱下载作业附件并生成 Excel；运行 python scripts/收作业.py --help。"""
import argparse
import getpass
import hashlib
import imaplib
import json
import os
import re
import ssl
import sys
import unicodedata
import zipfile
from datetime import date, timezone
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(PROJECT_ROOT / '本地依赖'))

from 学号范围 import in_scope
from 超大附件 import find_tasks, process_tasks, import_pdf, write_pending as write_large_pending


def parse_subject(subject):
    subject = unicodedata.normalize("NFKC", subject).strip()
    subject = re.sub(r'(第\s*[0-9一二三四五六七八九十百零〇两]+\s*周)高数作业', r'\1作业', subject)
    reverse = re.fullmatch(r'(第\s*[0-9一二三四五六七八九十百零〇两]+\s*周作业)\s+([^\d+\r\n]+?)\s+([0-9]{10})', subject)
    if reverse:
        subject = f'{reverse[1]}+{reverse[3]}+{reverse[2].strip()}'
    match = re.fullmatch(
        r"第\s*([0-9一二三四五六七八九十百零〇两]+)\s*周作业\s*[+_\-\s]*([0-9]{10})[+_\-\s]*([^+\r\n]+)", subject
    )
    if not match:
        return None
    number, sid, name = match.groups()
    if number.isascii() and number.isdigit():
        number = int(number)
    else:
        digits = dict(zip("零〇一二两三四五六七八九", [0, 0, 1, 2, 2, 3, 4, 5, 6, 7, 8, 9]))
        total = pending = 0
        for char in number:
            if char in digits:
                pending = digits[char]
            else:
                total += (pending or 1) * {"十": 10, "百": 100}[char]
                pending = 0
        number = total + pending
    name = name.strip()
    return (number, sid, name) if number > 0 and name else None


def safe_name(value):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', value).strip(' .')[:100]
    return value or "未命名"


def has_submission(record):
    """主题符合现有格式，且与记录的学生、周次一致，才计为已交。"""
    parsed = parse_subject(record.get('subject', ''))
    return bool(parsed and parsed[0] == record.get('number')
                and parsed[1] == record.get('sid'))


def legal_submission(record):
    """下载失败不取消提交资格；已确认无 PDF 的提交不能覆盖旧版本。"""
    if not record.get('sid') or not record.get('number') or not in_scope(record['sid']):
        return False
    if record.get('files') or record.get('pdf_attachment_names'):
        return True
    if 'pdf_attachment_names' not in record and any(
            str(name).lower().endswith('.pdf') for name in record.get('attachment_names', [])):
        return True
    return any(not any(reason in task.get('error', '') for reason in
                       ('不是 PDF', '缺少结束标记'))
               for task in record.get('large_attachments', []))


def submission_order(key, record):
    try:
        sent = parsedate_to_datetime(record.get('date', ''))
        stamp = (sent if sent.tzinfo else sent.replace(tzinfo=timezone.utc)).timestamp()
    except (TypeError, ValueError, OverflowError):
        stamp = 0
    uid = key.rsplit('|', 1)[-1]
    return stamp, int(uid) if uid.isdigit() else -1, key


def latest_submissions(state):
    chosen = {}
    for key, record in state.items():
        if not legal_submission(record):
            continue
        group = (record['sid'], record['number'])
        previous = chosen.get(group)
        if previous is None or submission_order(key, record) > submission_order(previous, state[previous]):
            chosen[group] = key
    return set(chosen.values())


def write_pending(root, state):
    selected = latest_submissions(state)
    write_large_pending(root, {key: record for key, record in state.items()
                              if key in selected or not legal_submission(record)})


def collection_issues(records):
    issues = []
    for record in records:
        reasons = [task.get('error') or '超大附件尚未下载'
                   for task in record.get('large_attachments', [])
                   if task.get('status') != '已下载']
        if record.get('collection_error'):
            reasons.append(record['collection_error'])
        if record.get('status') != '已交' and not reasons:
            reasons.append(record.get('status') or '收取状态未知')
        for reason in dict.fromkeys(reasons):
            issues.append([record.get('number'), record.get('sid', ''), record.get('name', ''),
                           '已交' if has_submission(record) else '待核对',
                           reason, record.get('subject', ''), record.get('date', '')])
    return issues


def subject_week(subject):
    """识别周次以限定收取范围；主题其他字段不规范时仍可记录本周问题。"""
    match = re.search(r'第\s*([0-9一二三四五六七八九十百零〇两]+)\s*周',
                      unicodedata.normalize('NFKC', subject))
    if match:
        parsed = parse_subject(f'第{match[1]}周作业+0000000000+待核对')
        return parsed[0] if parsed else None
    return None


def save_json(path, value):
    temporary = path.with_suffix('.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    temporary.replace(path)


def attachment_folder(root, record):
    return root / f"第{record['number']}周作业" / safe_name(f"{record['sid']}_{record['name']}")


def parse_since(value):
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        raise argparse.ArgumentTypeError('日期格式须为 YYYY-MM-DD，例如 2026-09-01')
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError('日期无效，请使用 YYYY-MM-DD，例如 2026-09-01') from exc


def collect(client, args, root, state):
    status, _ = client.select(args.mailbox, readonly=True)
    if status != 'OK':
        raise RuntimeError('无法打开邮箱文件夹')
    _, validity = client.response('UIDVALIDITY')
    if not validity or not validity[0]:
        raise RuntimeError('邮箱未返回 UIDVALIDITY，无法安全识别邮件')
    namespace = '|'.join([args.server, args.email, args.mailbox, validity[0].decode()])
    if args.since:
        months = ('Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec')
        since = f'{args.since.day:02d}-{months[args.since.month - 1]}-{args.since.year:04d}'
        status, data = client.uid('search', None, 'SINCE', since)
    else:
        status, data = client.uid('search', None, 'ALL')
    if status != 'OK':
        raise RuntimeError('搜索邮件失败')
    queue = []
    email_records = {}
    for uid in data[0].split():
        try:
            status, response = client.uid('fetch', uid, '(BODY.PEEK[HEADER.FIELDS (SUBJECT DATE)])')
            if status != 'OK':
                raise RuntimeError('读取邮件主题失败')
            headers = b''.join(item[1] for item in response if isinstance(item, tuple))
            if not headers:
                raise RuntimeError('邮件主题响应为空')
            header = BytesParser(policy=policy.default).parsebytes(headers)
            if args.assignment and subject_week(str(header.get('Subject', ''))) != args.assignment:
                continue
            queue.append((uid, {'date': str(header.get('Date', ''))}))
        except (OSError, ValueError, RuntimeError, imaplib.IMAP4.error) as exc:
            print(f'[失败] 邮件 UID {uid.decode()}：{exc}；下次运行会重试')
    for uid, _ in sorted(queue, key=lambda item: submission_order(namespace + '|' + item[0].decode(), item[1]), reverse=True):
        key = namespace + '|' + uid.decode()
        record = None
        try:
            status, response = client.uid('fetch', uid, '(BODY.PEEK[])')
            if status != 'OK':
                raise RuntimeError('读取邮件失败')
            raw = b''.join(item[1] for item in response if isinstance(item, tuple))
            if not raw:
                raise RuntimeError('邮件内容为空')
            msg = BytesParser(policy=policy.default).parsebytes(raw)
            subject = str(msg.get('Subject', ''))
            parsed = parse_subject(subject)
            if parsed and not in_scope(parsed[1]):
                continue
            if not parsed and '作业' not in subject:
                continue
            if parsed and args.assignment and parsed[0] != args.assignment:
                continue
            record = dict(subject=subject, sender=str(msg.get('From', '')),
                          date=str(msg.get('Date', '')), number=None, sid='', name='', files=[], status='主题格式错误')
            if parsed:
                number, sid, name = parsed
                record.update(number=number, sid=sid, name=name, status='无PDF附件')
                email_records[key] = record
                record['attachment_names'] = [part.get_filename() or '未命名附件'
                                              for part in msg.walk() if not part.is_multipart()
                                              and (part.get_filename() or part.get_content_disposition() == 'attachment')]
                record['pdf_attachment_names'] = []
                invalid_pdf = False
                for part in msg.walk():
                    filename = part.get_filename()
                    if part.is_multipart() or not (filename or part.get_content_disposition() == 'attachment'):
                        continue
                    if (filename and Path(filename).suffix.lower() == '.pdf') or (not filename and part.get_content_type() == 'application/pdf'):
                        payload = part.get_payload(decode=True) or b''
                        if payload.lstrip().startswith(b'%PDF-') and b'%%EOF' in payload[-2048:]:
                            record['pdf_attachment_names'].append(filename or '附件.pdf')
                        else:
                            invalid_pdf = True
                previous = state.get(key, {}).get('large_attachments', [])
                discovered = find_tasks(msg, key)
                tasks = {task['id']: dict(task) for task in previous
                         if task['url'] or task.get('file') or not any(t['url'] for t in discovered)}
                tasks.update({task['id']: next((dict(old) for old in previous if old['id'] == task['id']), task)
                              for task in discovered})
                record['large_attachments'] = list(tasks.values())
                if invalid_pdf:
                    record['collection_error'] = '存在损坏的 PDF 附件（文件头或结束标记无效）'
                candidates = {**state, key: record}
                if key not in latest_submissions(candidates):
                    # 旧附件保留为历史，不再下载，也不会加入本次作业包。
                    record['files'] = state.get(key, {}).get('files', [])
                    state[key] = record
                    save_json(root / '收取记录.json', state)
                    continue
                folder = attachment_folder(root, record)
                # 使用邮件唯一标识和附件序号区分重交及同名附件。
                token = hashlib.sha256(key.encode()).hexdigest()[:20]
                for index, part in enumerate(msg.walk(), 1):
                    if part.is_multipart():
                        continue
                    filename = part.get_filename()
                    if not filename and part.get_content_disposition() != 'attachment':
                        continue
                    if filename:
                        if Path(filename).suffix.lower() != '.pdf':
                            continue
                    elif part.get_content_type() == 'application/pdf':
                        filename = '附件.pdf'
                    else:
                        continue
                    payload = part.get_payload(decode=True)
                    if payload is None or not payload.lstrip().startswith(b'%PDF-') or b'%%EOF' not in payload[-2048:]:
                        continue
                    folder.mkdir(parents=True, exist_ok=True)
                    target = folder / f'{token}_{index}_{safe_name(filename or "附件.bin")}'
                    if not target.exists() or target.read_bytes() != payload:
                        temporary = target.with_suffix(target.suffix + '.tmp')
                        temporary.write_bytes(payload)
                        temporary.replace(target)
                    record['files'].append(str(target.relative_to(root)))
                process_tasks(root, record, folder, previous,
                              timeout=args.large_timeout, max_bytes=args.large_max_mb * 1024 ** 2)
            state[key] = record
            save_json(root / '收取记录.json', state)
            print(f"[{record['status']}] {subject}")
        except (OSError, ValueError, RuntimeError, imaplib.IMAP4.error) as exc:
            if record and record.get('sid') and record.get('number'):
                record.update(status='附件收取失败', collection_error=f'附件收取失败（{type(exc).__name__}），请重试或人工核对')
                record['files'] = list(dict.fromkeys(state.get(key, {}).get('files', []) + record['files']))
                state[key] = record
                save_json(root / '收取记录.json', state)
            print(f'[失败] 邮件 UID {uid.decode()}：{exc}；下次运行会重试')
    write_pending(root, state)
    from 整理学生邮箱 import organize
    organize(root, email_records)


def pack_homework(root, state, assignment=None):
    """按周平铺打包附件，文件名包含学生目录名以区分学生。"""
    root = Path(root).resolve()
    weeks = {assignment} if assignment else {r['number'] for r in state.values() if r['number']}
    if not assignment:
        weeks.update(int(m.group(1)) for p in root.iterdir()
                     if p.is_dir() and (m := re.fullmatch(r'第(\d+)周作业', p.name)))
    archives = []
    selected = latest_submissions(state)
    for week in sorted(weeks):
        folder = root / f'第{week}周作业'
        expected = {p for key, r in state.items() if key in selected and r['number'] == week
                    for p in r['files'] if Path(p).suffix.lower() == '.pdf'}
        missing = sum(not (root / p).is_file() for p in expected)
        if missing:
            print(f'[提醒] 第{week}周有 {missing} 个记录中的附件已不在本地，ZIP 仅包含当前存在的文件。')
        files = sorted(root / p for p in expected if (root / p).is_file())
        if not files:
            print(f'[提醒] 第{week}周没有可打包附件，将生成空 ZIP，避免保留旧版本。')
        for path in files:
            if not path.resolve().is_relative_to(folder.resolve()) or folder.is_symlink():
                raise ValueError('附件路径指向周作业目录以外，停止打包')
        destination = root / f'第{week}周作业.zip'
        temporary = destination.with_suffix('.zip.tmp')
        try:
            with zipfile.ZipFile(temporary, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
                used_names = set()
                for path in files:
                    name = '_'.join(path.relative_to(folder).parts)
                    candidate = name
                    suffix = 2
                    while candidate.casefold() in used_names:
                        candidate = f'{Path(name).stem}_{suffix}{Path(name).suffix}'
                        suffix += 1
                    used_names.add(candidate.casefold())
                    archive.write(path, candidate)
            temporary.replace(destination)
        finally:
            if temporary.exists():
                temporary.unlink()
        archives.append(destination)
        print(f'打包完成：{destination}（{len(files)} 个文件）')
    return archives


def read_roster(roster):
    if Path(roster).suffix.lower() == '.xls':
        import xlrd
        source = xlrd.open_workbook(str(roster))
        try:
            sheet = source.sheet_by_index(0)
            rows = [sheet.row_values(i) for i in range(sheet.nrows)]
        finally:
            source.release_resources()
    else:
        from openpyxl import load_workbook
        source = load_workbook(roster, read_only=True, data_only=True)
        try:
            rows = list(source.active.values)
        finally:
            source.close()
    header_index = next((i for i, row in enumerate(rows)
                         if {'学号', '姓名'} <= {str(v).strip() for v in row}), None)
    if header_index is None:
        raise ValueError('名单中未找到“学号”和“姓名”表头')
    headers = [str(v).strip() for v in rows[header_index]]
    id_col, name_col = headers.index('学号'), headers.index('姓名')
    students = {}
    for row in rows[header_index + 1:]:
        value = row[id_col]
        if value is None or str(value).strip() == '':
            continue
        sid = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()
        if not re.fullmatch(r'\d{10}', sid):
            raise ValueError(f'名单中的学号必须是10位：{sid}')
        if not in_scope(sid):
            continue
        name = str(row[name_col] or '').strip()
        if sid in students and students[sid] != name:
            raise ValueError(f'名单中同一学号对应多个姓名：{sid}')
        students[sid] = name
    return students


def export_excel(root, state, roster=None, assignment=None):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    records = [r for r in state.values() if not r['sid'] or in_scope(r['sid'])]
    selected = latest_submissions(state)
    students = read_roster(roster) if roster else {}
    roster_ids = set(students)
    previous_submitted = set()
    previous_notes = {}
    known_weeks = {r['number'] for r in records if r['number']} | ({assignment} if assignment else set())
    # 保留之前已建立的周列，包括已布置但尚无人提交的周次。
    existing_path = root / '作业统计.xlsx'
    if existing_path.exists():
        previous = load_workbook(existing_path, read_only=True, data_only=True)
        try:
            rows = iter(previous['提交统计'].values)
            headers = list(next(rows))
            week_columns = {}
            for index, value in enumerate(headers):
                match = re.fullmatch(r'第(\d+)周(?:作业)?', str(value))
                if match:
                    week = int(match.group(1))
                    known_weeks.add(week)
                    week_columns[index] = week
            for row in rows:
                value = row[headers.index('学号')]
                if value is None:
                    continue
                sid = str(value).strip()
                students.setdefault(sid, str(row[headers.index('姓名')] or '').strip())
                previous_submitted.update((sid, week) for col, week in week_columns.items() if row[col] == '已交')
                if '备注' in headers:
                    previous_notes[sid] = str(row[headers.index('备注')] or '')
        finally:
            previous.close()
    numbers = sorted(known_weeks)
    for r in records:
        if r['sid']:
            students.setdefault(r['sid'], r['name'])
    wb = Workbook()
    summary = wb.active
    summary.title = '提交统计'
    summary.append(['姓名', '学号'] + [f'第{n}周' for n in numbers] + ['已交周数', '备注'])
    submitted = previous_submitted | {(r['sid'], r['number']) for r in records if has_submission(r)}
    for sid, name in students.items():
        marks = ['已交' if (sid, n) in submitted else None for n in numbers]
        notes = [n for n in previous_notes.get(sid, '').split('；') if n]
        if roster:
            notes = [n for n in notes if n != '不在名单中']
        if roster and sid not in roster_ids:
            notes.append('不在名单中')
        if any(r['sid'] == sid and r['name'] != name for r in records):
            notes.append('邮件姓名与名单或其他提交不一致')
        summary.append([name, sid] + marks + [marks.count('已交'), '；'.join(dict.fromkeys(notes))])
    detail = wb.create_sheet('邮件明细')
    detail.append(['作业周次', '学号', '姓名', '状态', '邮件时间', '发件人', '主题', '附件数量', '附件相对路径', '超大附件待处理', '采用情况'])
    for r in records:
        pending = '\n'.join(f"{task['id']}：{task['error']}" for task in r.get('large_attachments', [])
                            if task['status'] != '已下载')
        chosen = any(state[key] is r for key in selected)
        disposition = '采用：最后一次合法提交' if chosen else ('历史提交：不采用' if legal_submission(r) else '不符合 PDF 收取要求')
        detail.append([r['number'], r['sid'], r['name'], r['status'], r['date'], r['sender'], r['subject'], len(r['files']), '\n'.join(r['files']), pending, disposition])
    issues = collection_issues([r for r in records if not legal_submission(r) or any(state[key] is r for key in selected)])
    errors = wb.create_sheet('收取异常')
    errors.append(['作业周次', '学号', '姓名', '提交认定', '异常原因', '邮件主题', '邮件时间'])
    for issue in issues:
        errors.append(issue)
    for ws in wb:
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='4472C4')
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(60, max(16, max(len(str(c.value or '')) for c in column) + 2))
    destination = root / '作业统计.xlsx'
    temporary = root / '作业统计.tmp.xlsx'
    wb.save(temporary)
    temporary.replace(destination)
    lines = ['作业收取异常记录（仅本地保存）',
             '主题格式正确的作业邮件计为已交；附件问题单独记录，不影响打钩。',
             '以下为当前待处理项；重新导出时更新，已解决项不再列出。', '']
    for week, sid, name, mark, reason, subject, sent_at in issues:
        lines.extend([f'第{week}周 | {sid} {name} | {mark}',
                      f'原因：{reason}', f'主题：{subject}', f'邮件时间：{sent_at}', ''])
    if not issues:
        lines.append('没有待处理的收取异常。')
    (root / '收取异常记录.txt').write_text('\n'.join(lines), encoding='utf-8-sig')
    print(f'统计已保存：{destination}')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--email', default=os.getenv('HOMEWORK_EMAIL'), help='收作业邮箱，也可设置 HOMEWORK_EMAIL')
    parser.add_argument('--server', default='imap.qq.com')
    parser.add_argument('--port', type=int, default=993)
    parser.add_argument('--mailbox', default='INBOX', help='默认收件箱')
    parser.add_argument('--output', default='作业收集')
    parser.add_argument('--roster', help='含学号、姓名的 .xls/.xlsx 名单；默认查找项目根目录下唯一的“选课名单”文件')
    parser.add_argument('--assignment', type=int, help='只下载指定周次；Excel 始终汇总本目录所有周次')
    parser.add_argument('--since', type=parse_since, help='仅下载此收件日期及之后的邮件附件，格式 YYYY-MM-DD（含当天）；不清除本地历史记录')
    parser.add_argument('--export-only', action='store_true', help='不连接邮箱，只根据本地收取记录重新生成 Excel')
    parser.add_argument('--pack-only', action='store_true', help='仅将本地周作业文件夹打包为 ZIP，不登录邮箱、不修改统计')
    parser.add_argument('--retry-large', action='store_true', help='仅重试本地记录中的超大附件，不登录邮箱；可配合 --assignment')
    parser.add_argument('--import-large', nargs=2, metavar=('任务编号', 'PDF路径'), help='补录网页下载的超大附件，不登录邮箱，并更新 ZIP、Excel 和网页')
    parser.add_argument('--large-timeout', type=int, default=120, help='超大附件单次网络读写超时秒数，默认120')
    parser.add_argument('--large-max-mb', type=int, default=10240, help='单个超大附件大小上限 MB，默认10240')
    parser.add_argument('--web-output', default='site', help='查询网页输出目录，默认 site，供 push 后发布')
    parser.add_argument('--title', default='高等数学 B', help='网页课程名称')
    args = parser.parse_args()
    if args.assignment is not None and args.assignment <= 0:
        parser.error('--assignment 必须大于0')
    if sum(bool(mode) for mode in (args.pack_only, args.export_only, args.retry_large, args.import_large)) > 1:
        parser.error('--pack-only、--export-only、--retry-large、--import-large 不能同时使用')
    if args.large_timeout <= 0 or args.large_max_mb <= 0:
        parser.error('--large-timeout 和 --large-max-mb 必须大于0')
    if args.import_large and (args.assignment or args.since):
        parser.error('--import-large 按任务编号定位，不与 --assignment 或 --since 同时使用')
    if args.retry_large and args.since:
        parser.error('--retry-large 使用本地任务，不与 --since 同时使用')
    try:
        if args.pack_only:
            root = Path(args.output).resolve()
            if not root.is_dir():
                raise ValueError('作业输出目录不存在')
            state_path = root / '收取记录.json'
            state = json.loads(state_path.read_text(encoding='utf-8')) if state_path.exists() else {}
            pack_homework(root, state, args.assignment)
            return
        from openpyxl import Workbook  # 下载之前检查依赖是否完整
        if not args.roster:
            candidates = [p for p in Path(__file__).resolve().parent.parent.glob('*选课名单*')
                          if p.suffix.lower() in ('.xls', '.xlsx') and not p.name.startswith('~$')]
            if len(candidates) > 1:
                raise ValueError('找到多份选课名单，请用 --roster 指定')
            if candidates:
                args.roster = candidates[0]
        if args.roster:
            read_roster(args.roster)  # 登录和下载前验证名单
        root = Path(args.output).resolve()
        root.mkdir(parents=True, exist_ok=True)
        state_path = root / '收取记录.json'
        state = json.loads(state_path.read_text(encoding='utf-8')) if state_path.exists() else {}
        if args.import_large:
            record = import_pdf(root, state, *args.import_large, attachment_folder)
            save_json(state_path, state)
            write_pending(root, state)
            pack_homework(root, state, record['number'])
            print(f"[补录完成] {record['subject']}")
        elif args.retry_large:
            for key in sorted(state, key=lambda k: submission_order(k, state[k]), reverse=True):
                record = state[key]
                if key not in latest_submissions(state):
                    continue
                if not record.get('large_attachments') or (args.assignment and record['number'] != args.assignment):
                    continue
                process_tasks(root, record, attachment_folder(root, record),
                              timeout=args.large_timeout, max_bytes=args.large_max_mb * 1024 ** 2)
                save_json(state_path, state)
            write_pending(root, state)
            pack_homework(root, state, args.assignment)
        elif not args.export_only:
            args.email = args.email or input('收作业邮箱：').strip()
            code = os.getenv('HOMEWORK_AUTH_CODE') or getpass.getpass('邮箱 IMAP 授权码（输入不显示）：')
            with imaplib.IMAP4_SSL(args.server, args.port, ssl_context=ssl.create_default_context(), timeout=60) as client:
                client.login(args.email, code)
                collect(client, args, root, state)
            pack_homework(root, state, args.assignment)
        export_excel(root, state, args.roster, args.assignment)
        from 生成网页 import build_site
        build_site(root / '作业统计.xlsx', args.web_output, args.title)
    except ImportError:
        parser.exit(1, '请先安装依赖：python -m pip install -r requirements.txt\n')
    except (OSError, ValueError, KeyError, StopIteration, imaplib.IMAP4.error) as exc:
        parser.exit(1, f'运行失败：{exc}\n若 Excel 正在打开，请关闭后重试；已保存的附件和记录会保留。\n')


if __name__ == '__main__':
    main()
