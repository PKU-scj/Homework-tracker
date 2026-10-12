"""离线验证 QQ 超大附件；所有邮件、HTTP 响应和下载文件均为合成数据。"""
import io
import hashlib
import json
import tempfile
import unittest
import zipfile
from email.message import EmailMessage
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import sys

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'scripts'))
import 超大附件 as large
from 收作业 import collect, import_pdf, attachment_folder, export_excel, pack_homework, parse_subject, latest_submissions


ENTRY = 'https://mail.qq.com/cgi-bin/ftnExs_download?k=test&code=abc'
DIRECT = 'https://download.ftn.qq.com/ftn_handler/test?fname=homework.pdf'
PDF = b'%PDF-1.4\n' + b'0' * (2 * large.CHUNK) + b'\n%%EOF\n'
MODERN = 'https://wx.mail.qq.com/ftn/download?func=3&key=test&code=test'
FILE_URL = 'https://wx.mail.qq.com/ftn/download?func=4&key=test&code=test'


def metadata(**changes):
    body = dict(name='作业.pdf', size=len(PDF), sha=hashlib.sha1(PDF).hexdigest(), url=FILE_URL)
    body.update(changes)
    return Response(json.dumps(dict(head={'ret': 0}, body=body)).encode(), MODERN, 'application/json')


class Response(io.BytesIO):
    def __init__(self, body, url=ENTRY, kind='application/pdf', length=None):
        super().__init__(body)
        self.status = 200
        self.url = url
        self.headers = EmailMessage()
        self.headers['Content-Type'] = kind
        self.headers['Content-Length'] = str(len(body) if length is None else length)

    def geturl(self):
        return self.url

    def read(self, size=-1):
        assert 0 <= size <= large.PAGE_LIMIT + 1, '必须有界读取，不能把超大文件整体读入内存'
        return super().read(size)


def message(body=None, ordinary=False, html=False):
    msg = EmailMessage()
    msg['Subject'] = '第3周作业+2600015447+示例甲'
    msg['From'] = 'student@example.com'
    msg.set_content(body if body is not None else '从QQ邮箱发来的超大附件\n' + ENTRY,
                    subtype='html' if html else 'plain')
    if ordinary:
        msg.add_attachment(b'%PDF-1.4\n%%EOF\n', maintype='application', subtype='pdf', filename='普通.pdf')
    return msg


class IMAP:
    def __init__(self, msg):
        self.msg = msg

    def select(self, mailbox, readonly):
        assert readonly
        return 'OK', []

    def response(self, name):
        return name, [b'42']

    def uid(self, action, *args):
        if action == 'search':
            return 'OK', [b'1']
        return 'OK', [(b'BODY[]', self.msg.as_bytes())]


class LargeAttachmentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.target = self.root / 'download.pdf'

    def test_latest_submission_by_date_then_numeric_uid(self):
        def record(day, **extra):
            return dict(number=4, sid='2600015447', name='示例甲',
                        date=f'Tue, {day:02d} Sep 2026 12:00:00 +0800',
                        files=[], pdf_attachment_names=['作业.pdf'], **extra)
        state = {'mail|10': record(2), 'mail|9': record(2), 'mail|11': record(1)}
        self.assertEqual(latest_submissions(state), {'mail|10'})
        state['empty|12'] = dict(record(3), pdf_attachment_names=[])
        state['zip|13'] = dict(record(4), pdf_attachment_names=[], large_attachments=[
            {'error': 'QQ 分享文件不是 PDF（可能为 ZIP 压缩包）'}])
        self.assertEqual(latest_submissions(state), {'mail|10'})
        state['failed|14'] = dict(record(5), pdf_attachment_names=[], large_attachments=[
            {'error': '下载连接失败'}])
        self.assertEqual(latest_submissions(state), {'failed|14'})
        state['other-week'] = dict(record(6), number=3)
        state['other-student'] = dict(record(6), sid='2600015448')
        self.assertEqual(latest_submissions(state), {'failed|14', 'other-week', 'other-student'})

    def test_collect_and_pack_only_latest_valid_email(self):
        messages = {}
        for uid, kind in ((1, 'pdf'), (2, 'pdf'), (3, 'empty'), (4, 'broken'), (5, 'zip'), (6, 'pdf')):
            msg = message('请查收', ordinary=kind == 'pdf')
            msg['Date'] = f'Tue, {(1 if uid == 6 else uid):02d} Sep 2026 12:00:00 +0800'
            if kind in ('broken', 'zip'):
                msg.add_attachment(b'invalid', maintype='application', subtype='octet-stream',
                                   filename='作业.pdf' if kind == 'broken' else '作业.zip')
            if uid == 2:
                msg.add_attachment(b'%PDF-1.4\n%%EOF\n', maintype='application', subtype='pdf', filename='第二页.pdf')
            messages[str(uid).encode()] = msg
        class ManyIMAP(IMAP):
            def uid(self, action, *args):
                if action == 'search':
                    return 'OK', [b'1 5 2 6 4 3']
                return 'OK', [(b'BODY[]', messages[args[0]].as_bytes())]
        args = SimpleNamespace(mailbox='INBOX', server='imap.qq.com', email='test@qq.com',
                               since=None, assignment=3, large_timeout=120, large_max_mb=10240)
        state = {}
        collect(ManyIMAP(None), args, self.root, state)
        self.assertEqual({k.rsplit('|', 1)[-1] for k in latest_submissions(state)}, {'2'})
        self.assertEqual(len(list(self.root.rglob('*.pdf'))), 2)
        # 不把历史残留文件或手动放入的无记录文件混入最终压缩包。
        orphan = self.root / '第3周作业' / '旧文件.pdf'
        orphan.write_bytes(PDF)
        archives = pack_homework(self.root, state, 3)
        with zipfile.ZipFile(archives[0]) as archive:
            self.assertEqual(len(archive.namelist()), 2)
            self.assertNotIn('旧文件.pdf', archive.namelist())
        chosen = state[next(iter(latest_submissions(state)))]
        chosen['files'] = []
        pack_homework(self.root, state, 3)
        with zipfile.ZipFile(archives[0]) as archive:
            self.assertEqual(archive.namelist(), [])

    def test_mime_html_entities_and_duplicates(self):
        msg = message()
        msg.add_alternative(f'<a href="{ENTRY.replace("&", "&amp;")}">下载</a>', subtype='html')
        tasks = large.find_tasks(msg, 'mail-key')
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0]['url'], ENTRY)
        self.assertEqual(tasks, large.find_tasks(msg, 'mail-key'))
        self.assertNotEqual(tasks[0]['id'], large.find_tasks(msg, 'different-mail')[0]['id'])

    def test_linkless_marker_is_pending_and_arbitrary_links_ignored(self):
        self.assertEqual(large.find_tasks(message('普通正文 https://example.com/a.pdf'), 'key'), [])
        tasks = large.find_tasks(message('超大附件 https://example.com/a.pdf'), 'key')
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0]['url'], '')

    def test_modern_qq_download_link_replaces_unknown_placeholder(self):
        state = {}
        self.run_collect(message('超大附件'), state)
        url = 'https://wx.mail.qq.com/ftn/download?func=3&k=test&key=test&code=test'
        with patch.object(large, 'build_opener') as opener:
            opener.return_value.open.side_effect = [metadata(), Response(PDF, FILE_URL)]
            record = self.run_collect(message('超大附件 ' + url), state)
        self.assertEqual(record['status'], '已交')
        self.assertEqual(len(record['large_attachments']), 1)
        self.assertRegex(Path(record['files'][0]).name, r'^[a-f0-9]{20}_1_作业\.pdf$')
        self.assertEqual(record['large_attachments'][0]['original_name'], '作业.pdf')
        self.assertFalse(large.allowed_url('https://wx.mail.qq.com/ftn/delete'))

    def test_modern_download_posts_share_parameters_and_checks_digest(self):
        with patch.object(large, 'build_opener') as opener:
            opener.return_value.open.side_effect = [metadata(), Response(PDF, FILE_URL)]
            large.download_pdf(MODERN, self.target)
            req = opener.return_value.open.call_args_list[0].args[0]
            self.assertEqual(req.get_method(), 'POST')
            self.assertIn(b'func=3', req.data)
            self.assertIn(b'f=json', req.data)
            self.assertNotIn('Cookie', req.headers)
        self.assertEqual(self.target.read_bytes(), PDF)
        self.assertTrue(large.allowed_url('https://gzc-dfsdown.mail.ftn.qq.com/1373/file?dkey=example'))
        self.assertFalse(large.allowed_url('https://gzc-dfsdown.mail.ftn.qq.com.evil.test/1373/file'))

    def test_modern_server_failure_and_unsafe_download_url(self):
        responses = [Response(b'{"head":{"ret":-1},"body":{}}', MODERN, 'application/json'),
                     metadata(url='http://127.0.0.1/private'), metadata(name='作业.zip'),
                     metadata(size=large.DEFAULT_LIMIT + 1), Response(b'not-json', MODERN)]
        for response in responses:
            with patch.object(large, 'build_opener') as opener:
                opener.return_value.open.return_value = response
                with self.assertRaises(large.DownloadProblem):
                    large.download_pdf(MODERN, self.target)
            self.assertFalse(self.target.exists())

    def test_modern_digest_mismatch_never_commits_file(self):
        with patch.object(large, 'build_opener') as opener:
            opener.return_value.open.side_effect = [metadata(sha='0' * 40), Response(PDF, FILE_URL)]
            with self.assertRaises(large.DownloadProblem):
                large.download_pdf(MODERN, self.target, attempts=1)
        self.assertFalse(self.target.exists())
        self.assertFalse(self.target.with_suffix('.pdf.part').exists())

    def test_modern_missing_credentials_and_cookie_injection(self):
        with patch.object(large, 'build_opener') as opener:
            with self.assertRaises(large.DownloadProblem):
                large.download_pdf('https://wx.mail.qq.com/ftn/download?func=3', self.target)
            opener.return_value.open.assert_not_called()
        with patch.object(large, 'build_opener') as opener:
            opener.return_value.open.return_value = metadata(cookie_name='token', cookie_value='a\r\nInjected: yes')
            with self.assertRaises(large.DownloadProblem):
                large.download_pdf(MODERN, self.target)
            self.assertEqual(opener.return_value.open.call_count, 1)

    def test_download_preserves_redirect_cookie_for_qq_file_server(self):
        from urllib.request import Request, HTTPCookieProcessor
        from http.cookiejar import CookieJar
        jar = CookieJar()
        processor = HTTPCookieProcessor(jar)
        response = Response(b'', FILE_URL)
        response.headers['Set-Cookie'] = 'mail5k=example-token; Domain=.qq.com; Path=/; Secure'
        response.info = lambda: response.headers
        processor.http_response(Request(FILE_URL), response)
        cdn = Request('https://gzc-dfsdown.mail.ftn.qq.com/1373/file')
        processor.http_request(cdn)
        self.assertEqual(cdn.get_header('Cookie'), 'mail5k=example-token')
        external = Request('https://example.com/')
        processor.http_request(external)
        self.assertIsNone(external.get_header('Cookie'))
        with patch.object(large, 'build_opener') as opener:
            opener.return_value.open.return_value = Response(PDF)
            large.download_pdf(ENTRY, self.target)
            handlers = opener.call_args.args
            self.assertTrue(any(isinstance(h, HTTPCookieProcessor) for h in handlers))
            self.assertTrue(any(isinstance(h, large.SafeRedirect) for h in handlers))

    def test_assignment_filters_other_weeks_before_reading_body(self):
        msg = message('超大附件')
        msg.replace_header('Subject', '第三周作业 姓名在前 2600015447')
        client = IMAP(msg)
        with patch.object(client, 'uid', wraps=client.uid) as uid:
            args = SimpleNamespace(mailbox='INBOX', server='imap.qq.com', email='test@qq.com',
                                   since=None, assignment=4, large_timeout=120, large_max_mb=10240)
            state = {}
            collect(client, args, self.root, state)
        self.assertEqual(state, {})
        self.assertFalse(any(call.args[-1] == '(BODY.PEEK[])' for call in uid.call_args_list))

    def test_alternative_fourth_week_subject_formats(self):
        for subject in ('第四周高数作业+ 2600015447 +示例甲', '第四周作业 示例甲 2600015447'):
            self.assertEqual(parse_subject(subject), (4, '2600015447', '示例甲'))

    def test_domains_and_redirects_are_restricted(self):
        for url in ('https://mail.qq.com.evil.test/cgi-bin/ftnExs_download',
                    'https://127.0.0.1/ftn_handler/a', 'file:///etc/passwd',
                    'https://mail.qq.com/cgi-bin/login',
                    'https://user:password@mail.qq.com/cgi-bin/ftnExs_download',
                    'https://mail.qq.com:8888/cgi-bin/ftnExs_download'):
            self.assertFalse(large.allowed_url(url), url)
        handler = large.SafeRedirect()
        with self.assertRaises(large.DownloadProblem):
            handler.redirect_request(None, None, 302, '', {}, 'http://127.0.0.1/private')

    def test_streamed_pdf_download(self):
        with patch.object(large, 'build_opener') as opener:
            opener.return_value.open.return_value = Response(PDF)
            large.download_pdf(ENTRY, self.target)
        self.assertEqual(self.target.read_bytes(), PDF)
        self.assertFalse(self.target.with_suffix('.pdf.part').exists())

    def test_landing_page_and_static_script_link(self):
        page = f'<script>location.href="{DIRECT.replace("/", chr(92) + "/")}";</script>'.encode()
        with patch.object(large, 'build_opener') as opener:
            opener.return_value.open.side_effect = [Response(page, kind='text/html'), Response(PDF, DIRECT)]
            large.download_pdf(ENTRY, self.target)
            self.assertEqual(opener.return_value.open.call_count, 2)
        self.assertEqual(self.target.read_bytes(), PDF)

    def test_login_page_and_non_pdf_not_saved(self):
        for body, kind in ((b'<html>Login required</html>', 'text/html'), (b'PK-zip', 'application/zip')):
            with patch.object(large, 'build_opener') as opener:
                opener.return_value.open.return_value = Response(body, kind=kind)
                with self.assertRaises(large.DownloadProblem):
                    large.download_pdf(ENTRY, self.target)
            self.assertFalse(self.target.exists())

    def test_truncated_transfer_retries_without_replacing_existing_file(self):
        self.target.write_bytes(b'old file')
        with patch.object(large, 'build_opener') as opener, patch.object(large.time, 'sleep'):
            opener.return_value.open.side_effect = [Response(PDF, length=len(PDF) + 10) for _ in range(3)]
            with self.assertRaises(large.DownloadProblem):
                large.download_pdf(ENTRY, self.target, progress=None)
            self.assertEqual(opener.return_value.open.call_count, 3)
        self.assertEqual(self.target.read_bytes(), b'old file')
        self.assertFalse(self.target.with_suffix('.pdf.part').exists())

    def test_transient_error_then_success(self):
        with patch.object(large, 'build_opener') as opener, patch.object(large.time, 'sleep'):
            opener.return_value.open.side_effect = [TimeoutError(), Response(PDF)]
            large.download_pdf(ENTRY, self.target, progress=None)
        self.assertTrue(large.valid_pdf(self.target))

    def test_size_limit_and_missing_eof(self):
        for body, limit in ((PDF, 1024), (b'%PDF-1.4\ntruncated', large.DEFAULT_LIMIT)):
            with patch.object(large, 'build_opener') as opener, patch.object(large.time, 'sleep'):
                opener.return_value.open.side_effect = [Response(body) for _ in range(3)]
                with self.assertRaises(large.DownloadProblem):
                    large.download_pdf(ENTRY, self.target, max_bytes=limit, progress=None)
            self.assertFalse(self.target.exists())

    def test_stream_without_length_still_enforces_size_limit(self):
        response = Response(PDF)
        del response.headers['Content-Length']
        with patch.object(large, 'build_opener') as opener:
            opener.return_value.open.return_value = response
            with self.assertRaises(large.DownloadProblem):
                large.download_pdf(ENTRY, self.target, max_bytes=large.CHUNK)
        self.assertFalse(self.target.exists())
        self.assertFalse(self.target.with_suffix('.pdf.part').exists())

    def run_collect(self, msg, state):
        args = SimpleNamespace(mailbox='INBOX', server='imap.qq.com', email='test@qq.com',
                               since=None, assignment=3, large_timeout=120, large_max_mb=10240)
        collect(IMAP(msg), args, self.root, state)
        return next(iter(state.values()))

    def test_collect_success_and_repeat_skips_network(self):
        state = {}
        with patch.object(large, 'build_opener') as opener:
            opener.return_value.open.return_value = Response(PDF)
            record = self.run_collect(message(ordinary=True), state)
            self.assertEqual(record['status'], '已交')
            self.assertEqual(len(record['files']), 2)
            self.run_collect(message(ordinary=True), state)
            self.assertEqual(opener.return_value.open.call_count, 1)
        self.assertEqual(len(next(iter(state.values()))['files']), 2)
        self.assertEqual(json.loads((self.root / '收取记录.json').read_text(encoding='utf-8')), state)

    def test_failed_task_import_and_repeat_preserve_submission(self):
        state = {}
        with patch.object(large, 'download_pdf', side_effect=large.DownloadProblem('需要登录')):
            record = self.run_collect(message(), state)
        self.assertEqual(record['status'], '超大附件待下载')
        task = record['large_attachments'][0]
        self.assertIn(task['id'], (self.root / '超大附件待处理.txt').read_text(encoding='utf-8-sig'))
        self.target.write_bytes(PDF)
        import_pdf(self.root, state, task['id'], self.target, attachment_folder)
        self.assertEqual(record['status'], '已交')
        self.assertTrue(record['files'][0].endswith('_1_download.pdf'))
        first_path = record['files'][0]
        import_pdf(self.root, state, task['id'], self.target, attachment_folder)
        self.assertEqual(record['files'], [first_path])
        with patch.object(large, 'download_pdf') as download:
            record = self.run_collect(message(), state)
            download.assert_not_called()
        self.assertEqual(record['status'], '已交')
        self.assertEqual(len(record['files']), 1)
        self.assertIn('没有待处理', (self.root / '超大附件待处理.txt').read_text(encoding='utf-8-sig'))

    def test_retry_missing_file_removes_stale_path_and_can_recover(self):
        state = {}
        with patch.object(large, 'build_opener') as opener:
            opener.return_value.open.return_value = Response(PDF)
            record = self.run_collect(message(), state)
        (self.root / record['files'][0]).unlink()
        with patch.object(large, 'download_pdf', side_effect=large.DownloadProblem('连接失败')):
            large.process_tasks(self.root, record, attachment_folder(self.root, record))
        self.assertEqual(record['status'], '超大附件待下载')
        self.assertEqual(record['files'], [])
        with patch.object(large, 'build_opener') as opener:
            opener.return_value.open.return_value = Response(PDF)
            large.process_tasks(self.root, record, attachment_folder(self.root, record))
        self.assertEqual(record['status'], '已交')
        self.assertEqual(len(record['files']), 1)

    def test_mixed_ordinary_pdf_and_failed_large_task(self):
        with patch.object(large, 'download_pdf', side_effect=large.DownloadProblem('需要登录')):
            record = self.run_collect(message(ordinary=True), {})
        self.assertEqual(record['status'], '已交')
        self.assertEqual(len(record['files']), 1)
        self.assertEqual(record['large_attachments'][0]['status'], '待下载')

    def test_ordinary_email_unchanged(self):
        with patch.object(large, 'download_pdf') as download:
            record = self.run_collect(message('普通正文', ordinary=True), {})
            download.assert_not_called()
        self.assertEqual(record['status'], '已交')
        self.assertEqual(record['large_attachments'], [])

    def test_non_pdf_attachment_counts_but_empty_email_does_not(self):
        from openpyxl import load_workbook
        for attached in (False, True):
            msg = message('请查收作业')
            if attached:
                msg.add_attachment(b'example', maintype='application', subtype='zip', filename='作业.zip')
            state = {}
            record = self.run_collect(msg, state)
            export_excel(self.root, state, assignment=3)
            wb = load_workbook(self.root / '作业统计.xlsx')
            self.assertEqual(wb['提交统计'].cell(2, 3).value, '已交' if attached else None)
            self.assertEqual(wb['收取异常'].cell(2, 4).value, '已交' if attached else '待核对')
            wb.close()

    def test_ordinary_attachment_write_failure_keeps_submission_evidence(self):
        from openpyxl import load_workbook
        state = {}
        with patch.object(Path, 'write_bytes', side_effect=OSError('disk failure')):
            record = self.run_collect(message('请查收', ordinary=True), state)
        self.assertEqual(record['status'], '附件收取失败')
        self.assertEqual(record['files'], [])
        export_excel(self.root, state, assignment=3)
        wb = load_workbook(self.root / '作业统计.xlsx')
        self.assertEqual(wb['提交统计'].cell(2, 3).value, '已交')
        self.assertIn('OSError', wb['收取异常'].cell(2, 5).value)
        wb.close()

    def test_import_rejects_invalid_pdf_and_unknown_task(self):
        state = {}
        record = self.run_collect(message('超大附件，无链接'), state)
        self.target.write_text('<html>login</html>')
        with self.assertRaises(ValueError):
            import_pdf(self.root, state, record['large_attachments'][0]['id'], self.target, attachment_folder)
        with self.assertRaises(ValueError):
            import_pdf(self.root, state, 'wrong-id', self.target, attachment_folder)
        self.assertEqual(record['status'], '超大附件待下载')

    def test_excel_zip_and_website_after_import(self):
        from openpyxl import load_workbook
        from 生成网页 import build_site
        state = {}
        record = self.run_collect(message('超大附件，无链接'), state)
        export_excel(self.root, state, assignment=3)
        wb = load_workbook(self.root / '作业统计.xlsx')
        self.assertEqual(wb['提交统计'].cell(2, 3).value, '已交')
        self.assertEqual(wb['收取异常'].cell(2, 4).value, '已交')
        self.assertIn('未识别到下载链接', (self.root / '收取异常记录.txt').read_text(encoding='utf-8-sig'))
        self.assertIn('未识别到下载链接', wb['邮件明细'].cell(2, 10).value)
        wb.close()
        site = build_site(self.root / '作业统计.xlsx', self.root / 'site')
        self.assertEqual(site['students'][0]['submitted'], [3])
        self.assertNotIn('未识别到下载链接', (self.root / 'site/index.html').read_text(encoding='utf-8'))
        self.target.write_bytes(PDF)
        import_pdf(self.root, state, record['large_attachments'][0]['id'], self.target, attachment_folder)
        export_excel(self.root, state, assignment=3)
        archives = pack_homework(self.root, state, 3)
        self.assertIn('没有待处理', (self.root / '收取异常记录.txt').read_text(encoding='utf-8-sig'))
        self.assertEqual(len(archives), 1)
        site = build_site(self.root / '作业统计.xlsx', self.root / 'site')
        self.assertEqual(site['students'][0]['submitted'], [3])
        self.assertNotIn('mail.qq.com', (self.root / 'site/index.html').read_text(encoding='utf-8'))


if __name__ == '__main__':
    unittest.main()
